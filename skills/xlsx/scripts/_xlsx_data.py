"""Bounded table reads and operation plans for the existing xlsx writer."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import tempfile
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from _xlsx_common import (
    EXCEL_INPUT_SUFFIXES, cell_range_bounds, input_file, output_file, publish_file,
    workbook_has_external_links,
)

MAX_DATA_CELLS = 500_000
SOURCE_ROW = "__source_row"


def require_columns(frame: pd.DataFrame, names: list[str]) -> list[str]:
    if not isinstance(names, list) or not names or len(set(names)) != len(names):
        raise ValueError("字段必须是非空、无重复的名称列表")
    missing = [name for name in names if name not in frame.columns]
    if missing:
        raise ValueError(f"字段不存在：{missing}；可选：{list(frame.columns)}")
    return names


def numeric(frame: pd.DataFrame, names: list[str], *, allow_missing: bool = False) -> pd.DataFrame:
    result = frame[require_columns(frame, names)].apply(pd.to_numeric, errors="raise")
    values = result.to_numpy(dtype=float, na_value=np.nan)
    if np.isinf(values).any() or (not allow_missing and np.isnan(values).any()):
        raise ValueError("数值字段含缺失值或无穷值；请先明确处理口径")
    return result.astype(float)


def scalar(value: Any) -> Any:
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if not math.isfinite(value):
            raise ValueError("结果含无穷值")
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_dataset(path_value: str, spec: dict[str, Any] | None = None) -> tuple[pd.DataFrame, dict]:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string, get_column_letter

    spec = spec or {}
    allowed = {"sheet", "range", "header_row", "columns", "exclude_rows", "numeric", "dates", "encoding", "path"}
    if set(spec) - allowed:
        raise ValueError(f"source 包含未知字段：{sorted(set(spec) - allowed)}")
    source = input_file(path_value, EXCEL_INPUT_SUFFIXES | {".csv", ".tsv"})
    if source.stat().st_size > 25 * 1024 * 1024:
        raise ValueError("分析输入上限为 25 MiB；请先分割文件")
    if workbook_has_external_links(source):
        raise ValueError("分析输入含外部链接，需先确认并固化其数据")
    header_row = int(spec.get("header_row", 1))
    if header_row < 1:
        raise ValueError("header_row 从 1 开始")
    bounds = cell_range_bounds(spec["range"]) if spec.get("range") else None
    if bounds and not bounds[1] <= header_row <= bounds[3]:
        raise ValueError("header_row 必须位于 range 内")
    records = []
    cached_wb = formula_wb = None
    formula_count = 0
    try:
        if source.suffix.lower() in EXCEL_INPUT_SUFFIXES:
            formula_wb = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            cached_wb = load_workbook(source, read_only=True, data_only=True, keep_links=False)
            active = formula_wb.active
            sheet_name = spec.get("sheet") or (active.title if active is not None else None)
            if sheet_name not in formula_wb.sheetnames:
                raise ValueError(f"工作表不存在：{sheet_name}")
            ws, cached = formula_wb[sheet_name], cached_wb[sheet_name]
            min_col, _, max_col, end_row = bounds or (1, header_row, ws.max_column, ws.max_row)
            if max_col < min_col or end_row < header_row or (end_row - header_row + 1) * (max_col - min_col + 1) > MAX_DATA_CELLS:
                raise ValueError("读取区域超过 500000 单元格或无效；请指定较小 range")
            rows = ws.iter_rows(min_row=header_row, max_row=end_row, min_col=min_col, max_col=max_col)
            values = cached.iter_rows(min_row=header_row, max_row=end_row, min_col=min_col, max_col=max_col)
            for row_number, (formula_row, cached_row) in enumerate(zip(rows, values), header_row):
                record = []
                for cell, cached_cell in zip(formula_row, cached_row):
                    is_formula = cell.data_type == "f"
                    value = cached_cell.value if is_formula else cell.value
                    if is_formula:
                        formula_count += 1
                        if value is None or cached_cell.data_type == "e":
                            raise ValueError(f"{sheet_name}!{cell.coordinate} 的公式缓存缺失或错误；先重算并核对")
                    if cell.data_type == "e":
                        raise ValueError(f"{sheet_name}!{cell.coordinate} 含 Excel 错误值")
                    record.append(value)
                records.append((row_number, record))
        else:
            if spec.get("sheet"):
                raise ValueError("CSV/TSV 没有工作表")
            sheet_name = None
            min_col, _, max_col, end_row = bounds or (1, header_row, 0, 1_048_576)
            with source.open(encoding=spec.get("encoding", "utf-8-sig"), newline="") as handle:
                reader = csv.reader(handle, delimiter="\t" if source.suffix.lower() == ".tsv" else ",")
                cells = 0
                for row_number, row in enumerate(reader, 1):
                    if row_number > end_row:
                        break
                    if row_number < header_row:
                        continue
                    max_col = max_col or len(row)
                    cells += max_col - min_col + 1
                    if cells > MAX_DATA_CELLS:
                        raise ValueError("读取区域超过 500000 单元格；请指定较小 range")
                    if not bounds and len(row) > max_col:
                        raise ValueError(f"CSV 第 {row_number} 行比表头多列，请先校正结构")
                    record = row[min_col - 1:max_col]
                    record += [None] * (max_col - min_col + 1 - len(record))
                    records.append((row_number, [None if v == "" else v for v in record]))
        if not records:
            raise ValueError("指定区域没有表头和数据")
        aliases = spec.get("columns")
        if aliases:
            if not isinstance(aliases, dict) or not all(isinstance(k, str) and k for k in aliases):
                raise ValueError("columns 必须为名称到 Excel 列字母的映射")
            names = list(aliases)
            offsets = [column_index_from_string(str(aliases[name]).upper()) - min_col for name in names]
            if len(set(offsets)) != len(offsets) or any(i < 0 or i > max_col - min_col for i in offsets):
                raise ValueError("columns 必须指向区域内不同的列")
        else:
            names = [str(v).strip() if v is not None else "" for v in records[0][1]]
            offsets = list(range(len(names)))
            if not all(names) or len(set(names)) != len(names):
                raise ValueError("表头为空或重复；请用 columns 显式指定名称到列字母的映射")
        if SOURCE_ROW in names:
            raise ValueError(f"{SOURCE_ROW} 为保留字段")
        excluded = set(spec.get("exclude_rows", []))
        if any(not isinstance(v, int) or v <= header_row for v in excluded):
            raise ValueError("exclude_rows 必须是表头之后的原始行号")
        data = [[row_number, *[row[i] for i in offsets]] for row_number, row in records[1:] if row_number not in excluded]
        frame = pd.DataFrame(data, columns=[SOURCE_ROW, *names], dtype=object)
        for name in spec.get("numeric", []):
            frame[name] = numeric(frame, [name], allow_missing=True)[name]
        for name, fmt in spec.get("dates", {}).items():
            require_columns(frame, [name])
            frame[name] = pd.to_datetime(frame[name], format=fmt, errors="raise")
        metadata = {
            "path": str(source), "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "sheet": sheet_name, "header_row": header_row,
            "columns": {name: get_column_letter(min_col + offset) for name, offset in zip(names, offsets)},
            "rows_read": len(records) - 1, "rows_used": len(frame),
            "excluded_rows": sorted(excluded), "formula_cache_cells": formula_count,
        }
        return frame, metadata
    finally:
        if cached_wb is not None:
            cached_wb.close()
        if formula_wb is not None:
            formula_wb.close()


def save_plan(tables: list[tuple[str, pd.DataFrame]], metadata: dict, destination: str,
              *, overwrite: bool = False, chart: dict | None = None) -> dict:
    from openpyxl.utils import get_column_letter

    target = output_file(destination, {".json"}, overwrite=overwrite)
    names = [name for name, _ in tables]
    if len(set(names)) != len(names) or "分析说明" in names:
        raise ValueError("结果工作表名称重复，或占用了保留名称 分析说明")
    note_rows = []

    def add_note(key: str, value: Any) -> None:
        if isinstance(value, dict) and value:
            for child, item in value.items():
                add_note(f"{key}.{child}", item)
            return
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=scalar)
        for start in range(0, max(1, len(text)), 80):
            note_rows.append([key if start == 0 else f"{key}（续）", text[start:start + 80]])

    for key, value in metadata.items():
        add_note(key, value)
    notes = pd.DataFrame(note_rows, columns=["项目", "说明"])
    tables = [*tables, ("分析说明", notes)]
    operations = []
    cell_count = 0
    for name, table in tables:
        if not name or len(name) > 31 or any(ch in name for ch in '[]:*?/\\'):
            raise ValueError(f"无效的工作表名称：{name}")
        if table.empty and len(table.columns) == 0:
            raise ValueError("结果表没有字段")
        rows = [[str(col) for col in table.columns], *[[scalar(v) for v in row] for row in table.itertuples(index=False, name=None)]]
        # Explicit value objects keep untrusted text out of Excel's formula parser.
        rows = [[{"value": v} if isinstance(v, str) else v for v in row] for row in rows]
        last_col, last_row = get_column_letter(len(table.columns)), len(rows)
        area = f"A1:{last_col}{last_row}"
        cell_count += len(table.columns) * len(rows) * 3 + len(table.columns)
        operations += [
            {"type": "add_sheet", "name": name},
            {"type": "write_rows", "sheet": name, "rows": rows},
            {"type": "style_range", "sheet": name, "range": area,
             "style": {"font": {"name": "Noto Sans CJK SC", "size": 11}, "alignment": {"vertical": "center"}}},
            {"type": "style_range", "sheet": name, "range": f"A1:{last_col}1",
             "style": {"font": {"bold": True, "color": "FFFFFF"}, "fill": {"color": "1F4E78"}}},
            {"type": "auto_fit", "sheet": name, "range": area, "min_width": 12},
            {"type": "freeze_panes", "sheet": name, "cell": "A2"},
            {"type": "set_print", "sheet": name, "print_area": area, "orientation": "landscape", "fit_to_width": 1, "fit_to_height": 0, "repeat_rows": "1:1"},
        ]
        for column_index, column in enumerate(table.columns, 1):
            values = table[column].dropna()
            if len(values) and all(isinstance(v, (int, float, np.number)) and not isinstance(v, (bool, np.bool_)) for v in values):
                letter = get_column_letter(column_index)
                number_format = "#,##0" if all(float(v).is_integer() for v in values) else "#,##0.####"
                operations.append({"type": "style_range", "sheet": name, "range": f"{letter}2:{letter}{last_row}",
                                   "style": {"number_format": number_format, "alignment": {"horizontal": "right", "indent": 1}}})
                cell_count += len(table)
    if cell_count > 100_000:
        raise ValueError("结果超过 apply_workbook 的单次处理上限；请缩小结果范围或分组处理")
    if chart:
        name, table = tables[0]
        category = chart["category"]
        values = chart["values"]
        require_columns(table, [category, *values])
        if not table.columns.is_unique:
            raise ValueError("图表来源包含重复字段名")
        column_names = list(table.columns)
        indexes = [column_names.index(value) + 1 for value in values]
        if indexes != list(range(min(indexes), max(indexes) + 1)):
            raise ValueError("图表 values 需按顺序选择相邻的结果列")
        c = get_column_letter(column_names.index(category) + 1)
        end = len(table) + 1
        if end < 2:
            raise ValueError("没有数据可用于图表")
        operations += [{"type": "add_chart", "sheet": name, "chart_type": chart.get("type", "column"),
                        "title": chart.get("title", name), "data_range": f"{get_column_letter(min(indexes))}1:{get_column_letter(max(indexes))}{end}",
                        "categories_range": f"{c}2:{c}{end}", "anchor": f"A{end+3}", "width": 20, "height": 10,
                        "legend_position": "b", "x_axis_title": category, "y_axis_title": chart.get("value_title", "数值")},
                       {"type": "set_print", "sheet": name, "print_area": f"A1:{get_column_letter(max(16, len(table.columns)))}{end+27}", "orientation": "landscape", "fit_to_width": 1, "fit_to_height": 0}]
    plan = {"operations": operations, "active_sheet": tables[0][0]}
    raw = json.dumps(plan, ensure_ascii=False, allow_nan=False, default=scalar).encode()
    if len(raw) > 2 * 1024 * 1024:
        raise ValueError("结果操作说明超过 2 MiB；请缩小结果或分批输出")
    descriptor, temp_name = tempfile.mkstemp(dir=target.parent, suffix=".json")
    temp = Path(temp_name)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(raw)
        publish_file(temp, target, overwrite=overwrite)
    finally:
        temp.unlink(missing_ok=True)
    return {"spec_path": str(target), "tables": [{"sheet": n, "rows": len(t), "columns": len(t.columns)} for n, t in tables],
            "next_script": "scripts/apply_workbook.py", "metadata": metadata}
