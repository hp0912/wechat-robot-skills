import csv
import hashlib
import sys
import unittest
import tempfile
from copy import copy
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import _xlsx_common as common
import _xlsx_data as data
import analyze_workbook as analysis
import apply_workbook as writer
import inspect_workbook as inspector
import model_workbook as modeling

ROOT = Path(__file__).resolve().parents[1]

_TEMP = tempfile.TemporaryDirectory(prefix="xlsx-tests-")
QA = Path(_TEMP.name).resolve()
_ORIGINAL_OUTPUT_ROOT = common.EXCEL_OUTPUT_ROOT


def tearDownModule():
    common.EXCEL_OUTPUT_ROOT = _ORIGINAL_OUTPUT_ROOT
    _TEMP.cleanup()


class MergeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        common.EXCEL_OUTPUT_ROOT = QA
        cls.source = QA / "sales.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "明细"
        for row in [
            ["地区", "收入", "成本", "说明"],
            ["华东", 100, 60, "满意"],
            ["华东", -20, 5, "退款"],
            ["华南", 80, 50, "物流慢"],
            ["华南", None, 20, "服务好但物流慢"],
            ["合计", 160, 135, None],
        ]:
            ws.append(row)
        font = copy(ws["A1"].font)
        font.bold = True
        ws["A1"].font = font
        wb.save(cls.source)
        cls.source_hash = hashlib.sha256(cls.source.read_bytes()).hexdigest()
        cls.csv = QA / "text.csv"
        with cls.csv.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(
                [["ID", "文本"], ["001", "=1+1"], ["002", "长文本" * 80]]
            )

    def dataset(self):
        return data.read_dataset(
            str(self.source), {"sheet": "明细", "exclude_rows": [6]}
        )[0]

    def test_profile_finds_totals_without_dropping_negative_rows(self):
        _, result = analysis.analyze(str(self.source), {"method": "profile"})
        self.assertEqual(result["source"]["rows_used"], 5)
        self.assertEqual(result["summary_row_candidates"][0]["row"], 6)

    def test_aggregate_negative_values_and_null_count(self):
        tables, _ = analysis.analyze(
            str(self.source),
            {
                "method": "aggregate",
                "source": {"exclude_rows": [6]},
                "by": ["地区"],
                "metrics": {"收入": "sum", "说明": "count"},
            },
        )
        result = tables[0][1].set_index("地区")
        self.assertEqual(result.loc["华东", "收入"], 80)
        self.assertEqual(result.loc["华南", "说明"], 2)

    def test_all_null_sum_not_zero(self):
        frame = pd.DataFrame({"类别": ["A", "B", "B"], "值": [None, 0, None]})
        result = analysis.aggregate(
            frame, {"by": ["类别"], "metrics": {"值": "sum"}}, pivot=False
        ).set_index("类别")
        self.assertTrue(pd.isna(result.loc["A", "值"]))
        self.assertEqual(result.loc["B", "值"], 0)

    def test_pivot_multiple_levels(self):
        frame = pd.DataFrame(
            {
                "地区": ["A", "A", "B"],
                "年": ["2025", "2026", "2025"],
                "收入": [10, 20, 30],
            }
        )
        result = analysis.aggregate(
            frame,
            {"by": ["地区"], "columns": ["年"], "metrics": {"收入": "sum"}},
            pivot=True,
        ).set_index("地区")
        self.assertEqual(result.loc["A", '["收入","2026"]'], 20)
        self.assertTrue(pd.isna(result.loc["B", '["收入","2026"]']))

    def test_rules_report_conflict_and_unknown(self):
        detail, summary = analysis.classify(
            self.dataset(),
            {
                "column": "说明",
                "rules": [
                    {"label": "正向", "keywords": ["好", "满意"]},
                    {"label": "物流", "keywords": ["慢"]},
                ],
            },
        )
        self.assertEqual(detail["分类"].tolist(), ["正向", "未知", "物流", "需复核"])
        self.assertAlmostEqual(summary["占比"].sum(), 1)

    def test_join_rejects_accidental_many_to_many(self):
        with self.assertRaises(pd.errors.MergeError):
            analysis.transform(
                self.dataset(),
                [
                    {
                        "type": "merge",
                        "source": {"path": str(self.source), "exclude_rows": [6]},
                        "on": ["地区"],
                    }
                ],
                [],
            )

    def test_missing_headers_can_use_coordinates(self):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append([None, "值"])
        ws.append(["001", 5])
        path = QA / "no_header.xlsx"
        wb.save(path)
        with self.assertRaises(ValueError):
            data.read_dataset(str(path))
        frame, meta = data.read_dataset(
            str(path), {"columns": {"编号": "A", "值": "B"}}
        )
        self.assertEqual(frame.iloc[0]["编号"], "001")
        self.assertEqual(meta["columns"]["编号"], "A")

    def test_formula_cache_is_required(self):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["值"])
        ws.append(["=1+1"])
        path = QA / "uncached.xlsx"
        wb.save(path)
        with self.assertRaisesRegex(ValueError, "公式缓存"):
            data.read_dataset(str(path))

    def test_safe_text_and_no_truncation_in_writer(self):
        tables, metadata = analysis.analyze(str(self.csv), {"method": "transform"})
        plan = QA / "safe-text.json"
        output = QA / "safe-text.xlsx"
        data.save_plan(tables, metadata, str(plan), overwrite=True)
        with patch.object(
            sys,
            "argv",
            ["apply", "--output", str(output), "--spec-file", str(plan), "--overwrite"],
        ):
            result = writer.main()
        self.assertEqual(result["formula_count"], 0)
        wb = load_workbook(output)
        ws = wb["分析结果"]
        self.assertEqual(ws["B2"].value, "001")
        self.assertEqual(ws["C2"].value, "=1+1")
        self.assertEqual(ws["C2"].data_type, "s")
        self.assertEqual(ws["C3"].value, "长文本" * 80)
        self.assertTrue(ws["C3"].alignment.wrap_text)
        self.assertGreater(ws.row_dimensions[3].height, 36)
        self.assertEqual(
            inspector.inspect_excel(
                output,
                sheet_name="分析结果",
                start_row=1,
                start_column=1,
                max_rows=10,
                max_columns=10,
            )["formula_count"],
            0,
        )

    def test_append_result_preserves_source(self):
        spec = {
            "method": "aggregate",
            "source": {"exclude_rows": [6]},
            "by": ["地区"],
            "metrics": {"收入": "sum"},
            "chart": {"category": "地区", "values": ["收入"], "title": "地区收入"},
        }
        tables, meta = analysis.analyze(str(self.source), spec)
        plan = QA / "summary.json"
        output = QA / "summary.xlsx"
        data.save_plan(tables, meta, str(plan), overwrite=True, chart=spec["chart"])
        with patch.object(
            sys,
            "argv",
            [
                "apply",
                "--input",
                str(self.source),
                "--output",
                str(output),
                "--spec-file",
                str(plan),
                "--overwrite",
            ],
        ):
            writer.main()
        self.assertEqual(
            hashlib.sha256(self.source.read_bytes()).hexdigest(), self.source_hash
        )
        original = load_workbook(self.source)
        result = load_workbook(output)
        self.assertEqual(list(original["明细"].values), list(result["明细"].values))
        self.assertEqual(
            copy(original["明细"]["A1"].font), copy(result["明细"]["A1"].font)
        )
        self.assertEqual(len(getattr(result["分析结果"], "_charts")), 1)

    def test_writer_rejects_text_that_excel_would_truncate(self):
        for value in ["长" * 32768, {"value": "长" * 32768}]:
            with self.subTest(explicit=isinstance(value, dict)):
                wb = Workbook()
                assert wb.active is not None
                with self.assertRaisesRegex(ValueError, "不能静默截断"):
                    writer._op_write_rows(
                        wb, {"sheet": wb.active.title, "rows": [[value]]}
                    )

    def test_cost_direction_once(self):
        frame = pd.DataFrame(
            {
                data.SOURCE_ROW: [2, 3],
                "对象": ["便宜", "贵"],
                "质量": [98, 98],
                "价格": [10, 30],
            }
        )
        tables, _ = modeling.evaluate(
            frame, {"entity": "对象", "directions": {"质量": "benefit", "价格": "cost"}}
        )
        rank = tables[0][1].set_index("对象")
        self.assertEqual(rank.loc["便宜", "排名"], 1)
        self.assertEqual(rank.loc["贵", "排名"], 2)

    def test_identical_objects_tie(self):
        frame = pd.DataFrame(
            {data.SOURCE_ROW: [2, 3], "对象": ["A", "B"], "值": [10, 10]}
        )
        tables, _ = modeling.evaluate(
            frame,
            {"entity": "对象", "directions": {"值": "cost"}, "weighting": "entropy"},
        )
        self.assertEqual(tables[0][1]["排名"].tolist(), [1, 1])
        self.assertEqual(tables[0][1]["得分"].tolist(), [0.5, 0.5])

    def test_ahp_rejects_inconsistent_matrix(self):
        frame = pd.DataFrame(
            {
                data.SOURCE_ROW: [2, 3],
                "对象": ["A", "B"],
                "a": [1, 2],
                "b": [2, 1],
                "c": [2, 3],
            }
        )
        with self.assertRaisesRegex(ValueError, "一致性"):
            modeling.evaluate(
                frame,
                {
                    "entity": "对象",
                    "directions": {"a": "benefit", "b": "benefit", "c": "benefit"},
                    "weighting": "ahp",
                    "comparison_matrix": [[1, 9, 1 / 9], [1 / 9, 1, 9], [9, 1 / 9, 1]],
                },
            )

    def test_forecast_by_group_and_time_holdout(self):
        dates = pd.date_range("2024-01-01", periods=24, freq="MS")
        frame = pd.DataFrame(
            {
                "城市": ["A"] * 24 + ["B"] * 24,
                "月份": list(dates) * 2,
                "销量": list(np.arange(24) * 10 + 100) + list(np.arange(24) * -2 + 100),
            }
        )
        tables, _ = modeling.forecast(
            frame,
            {
                "by": ["城市"],
                "date": "月份",
                "value": "销量",
                "horizon": 3,
                "frequency": "MS",
            },
        )
        result = tables[0][1]
        self.assertEqual(len(result), 6)
        self.assertAlmostEqual(result[result["城市"] == "A"].iloc[0]["预测值"], 340)
        self.assertAlmostEqual(result[result["城市"] == "B"].iloc[0]["预测值"], 52)
        self.assertTrue((tables[1][1]["训练期数"] == 18).all())

    def test_forecast_rejects_missing_month(self):
        frame = pd.DataFrame(
            {
                "日期": pd.date_range("2024-01-01", periods=8, freq="MS").delete(3),
                "值": range(7),
            }
        )
        with self.assertRaisesRegex(ValueError, "连续"):
            modeling.forecast(frame, {"date": "日期", "value": "值"})

    def test_regression_has_holdout_and_baseline(self):
        frame = pd.DataFrame(
            {data.SOURCE_ROW: range(2, 42), "x": range(40), "y": np.arange(40) * 3 + 7}
        )
        tables, meta = modeling.supervised(
            frame, {"features": ["x"], "target": "y"}, classification=False
        )
        metrics = tables[1][1]
        error = metrics[
            (metrics["数据集"] == "测试")
            & (metrics["模型"] == "linear")
            & (metrics["指标"] == "RMSE")
        ].iloc[0]["值"]
        self.assertLess(error, 1e-8)
        self.assertEqual(meta["train_rows"], 32)
        self.assertIn("基线", metrics["模型"].tolist())

    def test_classification_categorical_pipeline(self):
        frame = pd.DataFrame(
            {
                data.SOURCE_ROW: range(2, 42),
                "x": list(range(20)) * 2,
                "组": ["A"] * 20 + ["B"] * 20,
                "标签": ["低"] * 20 + ["高"] * 20,
            }
        )
        tables, _ = modeling.supervised(
            frame,
            {"features": ["x", "组"], "categorical": ["组"], "target": "标签"},
            classification=True,
        )
        self.assertEqual(len(tables[0][1]), 8)
        self.assertIn("F1_macro", tables[1][1]["指标"].tolist())

    def test_small_regression_has_two_test_rows_for_r_squared(self):
        frame = pd.DataFrame(
            {data.SOURCE_ROW: range(2, 12), "x": range(10), "y": np.arange(10) * 3 + 7}
        )
        tables, metadata = modeling.supervised(
            frame,
            {"features": ["x"], "target": "y", "test_fraction": 0.1},
            classification=False,
        )
        self.assertEqual(metadata["test_rows"], 2)
        self.assertTrue(np.isfinite(tables[1][1]["值"]).all())

    def test_clustering_separates_obvious_groups(self):
        frame = pd.DataFrame(
            {data.SOURCE_ROW: range(8), "x": [0, 0.1, 0.2, 0.3, 10, 10.1, 10.2, 10.3]}
        )
        tables, _ = modeling.unsupervised(
            frame, {"features": ["x"], "clusters": 2}, anomaly=False
        )
        labels = tables[0][1]["簇编号"].to_numpy()
        self.assertTrue(np.all(labels[:4] == labels[0]))
        self.assertNotEqual(labels[0], labels[-1])

    def test_integer_optimization(self):
        tables, meta = modeling.optimize(
            {
                "variables": ["x", "y"],
                "objective": [3, 2],
                "sense": "max",
                "integer": [True, True],
                "constraints": [{"coefficients": [2, 1], "relation": "<=", "rhs": 4}],
            }
        )
        self.assertEqual(meta["objective_value"], 8)
        self.assertEqual(tables[0][1]["取值"].tolist(), [0, 4])

    def test_infeasible_and_unbounded_rejected(self):
        with self.assertRaisesRegex(ValueError, "求解未成功"):
            modeling.optimize(
                {
                    "variables": ["x"],
                    "objective": [1],
                    "constraints": [{"coefficients": [1], "relation": "<=", "rhs": -1}],
                }
            )
        with self.assertRaisesRegex(ValueError, "求解未成功"):
            modeling.optimize({"variables": ["x"], "objective": [1], "sense": "max"})

    def test_convex_quadratic(self):
        tables, meta = modeling.optimize(
            {"variables": ["x"], "objective": [-4], "quadratic": [[2]]}
        )
        self.assertAlmostEqual(tables[0][1].iloc[0]["取值"], 2)
        self.assertAlmostEqual(meta["objective_value"], -4)

    def test_output_root_enforced(self):
        with self.assertRaises(ValueError):
            data.save_plan(
                [("结果", pd.DataFrame({"x": [1]}))],
                {},
                "/private/tmp/outside-plan.json",
            )


if __name__ == "__main__":
    unittest.main(verbosity=2)
